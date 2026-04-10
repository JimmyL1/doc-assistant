from pathlib import Path

from langchain_core.documents import Document


def parse_xlsx(file_path: str | Path) -> list[Document]:
    file_path = Path(file_path)

    if file_path.suffix.lower() == ".xls":
        return _parse_xls(file_path)
    return _parse_xlsx(file_path)


def _parse_xlsx(file_path: Path) -> list[Document]:
    import openpyxl

    wb = openpyxl.load_workbook(file_path, data_only=True)
    documents: list[Document] = []

    for sheet_index, sheet in enumerate(wb.worksheets):
        rows = list(sheet.iter_rows(values_only=True))
        doc = _rows_to_document(rows, file_path.name, sheet_index, sheet.title)
        if doc:
            documents.append(doc)

    if not documents:
        raise ValueError(f"No text could be extracted from '{file_path.name}'")
    return documents


def _parse_xls(file_path: Path) -> list[Document]:
    import xlrd

    wb = xlrd.open_workbook(str(file_path))
    documents: list[Document] = []

    for sheet_index in range(wb.nsheets):
        sheet = wb.sheet_by_index(sheet_index)
        rows = [tuple(sheet.row_values(r)) for r in range(sheet.nrows)]
        doc = _rows_to_document(rows, file_path.name, sheet_index, sheet.name)
        if doc:
            documents.append(doc)

    if not documents:
        raise ValueError(f"No text could be extracted from '{file_path.name}'")
    return documents


def _rows_to_document(
    rows: list[tuple], source: str, sheet_index: int, sheet_name: str
) -> Document | None:
    if not rows:
        return None

    headers = [str(cell) if cell is not None else "" for cell in rows[0]]
    lines: list[str] = []

    for row in rows[1:]:
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        parts = [
            f"{headers[i]}: {str(cell) if cell is not None else ''}"
            for i, cell in enumerate(row)
            if i < len(headers)
        ]
        lines.append(" | ".join(parts))

    if not lines:
        return None

    return Document(
        page_content="\n".join(lines),
        metadata={"source": source, "file_type": "xlsx", "page": sheet_index, "sheet": sheet_name},
    )
